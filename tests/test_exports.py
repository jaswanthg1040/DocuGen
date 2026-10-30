import zipfile
from io import BytesIO

import pytest
from docx import Document

from services.formatting import parse_blocks, preview_markdown
from services.pdf_export import build_pdf
from services.utils import safe_filename
from services.word_export import build_docx

SAMPLE = """# Main heading

First paragraph with **bold**, *italic* and ***both***.
Second line in the same paragraph.

## Sub heading
- bullet one
- bullet **two**

1. step one
2. step two

Text between lists.

1. restart at one
"""


def pdf_pages(data: bytes) -> int:
    return data.count(b"/Type /Page") - data.count(b"/Type /Pages")


def test_docx_is_valid_zip_and_reopens():
    data = build_docx("My Doc", SAMPLE)
    assert zipfile.is_zipfile(BytesIO(data))
    doc = Document(BytesIO(data))
    assert doc.paragraphs[0].text == "My Doc"
    assert doc.core_properties.title == "My Doc"


def test_docx_preserves_structure_and_formatting():
    doc = Document(BytesIO(build_docx("T", SAMPLE, "A description")))
    styles = [(p.style.name, p.text) for p in doc.paragraphs]
    assert ("Heading 1", "Main heading") in styles
    assert ("Heading 2", "Sub heading") in styles
    assert [t for s, t in styles if s == "List Bullet"] == ["bullet one", "bullet two"]
    numbered = [t for s, t in styles if t[:2] in ("1.", "2.")]
    assert numbered == ["1.\tstep one", "2.\tstep two", "1.\trestart at one"]
    assert any(t == "A description" for _, t in styles)
    para = next(p for p in doc.paragraphs if p.text.startswith("First paragraph"))
    runs = {r.text: (bool(r.bold), bool(r.italic)) for r in para.runs if r.text}
    assert runs["bold"] == (True, False)
    assert runs["italic"] == (False, True)
    assert runs["both"] == (True, True)
    assert "Second line in the same paragraph." in para.text and "\n" in para.text  # line break kept


def test_docx_empty_document_and_empty_title():
    doc = Document(BytesIO(build_docx("", "")))
    assert doc.paragraphs[0].text == "Untitled document"
    assert "empty" in doc.paragraphs[1].text.lower()


def test_docx_description_toggle_and_font():
    with_desc = Document(BytesIO(build_docx("T", "x", "Hello desc", "Arial", True)))
    without = Document(BytesIO(build_docx("T", "x", "Hello desc", "Arial", False)))
    assert any(p.text == "Hello desc" for p in with_desc.paragraphs)
    assert not any(p.text == "Hello desc" for p in without.paragraphs)
    assert with_desc.styles["Normal"].font.name == "Arial"


def test_docx_unicode_and_control_characters():
    text = "Caf\u00e9 \u00fcber \u4f60\u597d \u0c24\u0c46\u0c32\u0c41\u0c17\u0c41 \U0001F600\x00\x0b end"
    doc = Document(BytesIO(build_docx("T\u00edtulo \u2713\x00", text)))
    assert "Caf\u00e9" in "\n".join(p.text for p in doc.paragraphs)


def test_pdf_is_valid():
    data = build_pdf("My Doc", SAMPLE, "desc")
    assert data.startswith(b"%PDF-") and data.rstrip().endswith(b"%%EOF")


def test_pdf_empty_and_special_title():
    assert build_pdf("", "").startswith(b"%PDF-")
    assert build_pdf("<b>&\"'/\\:*?", "text <tag> & more").startswith(b"%PDF-")


def test_pdf_long_document_has_multiple_pages():
    body = "\n\n".join(f"Paragraph {i}. " + "word " * 120 for i in range(80))
    data = build_pdf("Long", body)
    assert pdf_pages(data) > 3


def test_pdf_handles_very_long_unbroken_word():
    assert build_pdf("Long word", "x" * 5000 + "\n\n" + "y" * 300).startswith(b"%PDF-")


def test_pdf_unicode_content():
    data = build_pdf("Unicode \u2713", "Caf\u00e9 \u00fcber \u201cquotes\u201d \u2014 \u20ac \u4f60\u597d \U0001F600")
    assert data.startswith(b"%PDF-")


def test_pdf_lists_and_formatting_do_not_crash():
    assert build_pdf("T", SAMPLE, "d", include_description=False).startswith(b"%PDF-")


@pytest.mark.parametrize("title,expected", [
    ("Report 2024", "Report_2024.docx"),
    ('a/b\\c:d*e?f"g<h>i|j', "a_b_c_d_e_f_g_h_i_j.docx"),
    ("", "document.docx"),
    ("...", "document.docx"),
    ("CON", "document.docx"),
    ("../../etc/passwd", "etc_passwd.docx"),
])
def test_safe_filename(title, expected):
    assert safe_filename(title, "docx") == expected


def test_safe_filename_length_and_unicode():
    assert len(safe_filename("x" * 500, "pdf")) <= 84
    assert safe_filename("\u00dcber caf\u00e9", "pdf") == "\u00dcber_caf\u00e9.pdf"


def test_parser_blocks():
    kinds = [(b.kind, b.level or b.number) for b in parse_blocks(SAMPLE)]
    assert kinds == [("heading", 1), ("paragraph", 0), ("heading", 2), ("bullet", 0), ("bullet", 0),
                     ("number", 1), ("number", 2), ("paragraph", 0), ("number", 1)]
    assert parse_blocks("") == [] and parse_blocks(None) == []
    assert parse_blocks("#nospace")[0].kind == "paragraph"
    assert parse_blocks("2 * 3 * 4")[0].lines[0] == [("2 * 3 * 4", False, False)]


def test_preview_escapes_unsupported_markup():
    md = preview_markdown("<script>alert(1)</script> [link](http://x) ![img](http://y) $x$ **ok**")
    assert "<script>" not in md.replace("\\<", "")
    assert "\\[link\\]" in md and "\\!" not in md or "\\[img\\]" in md
    assert "**ok**" in md
    assert preview_markdown("") == ""
