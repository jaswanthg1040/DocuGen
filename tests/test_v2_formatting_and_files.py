import zipfile
from io import BytesIO

import pytest
from docx import Document

from services.formatting import count_stats, parse_blocks, preview_markdown
from services.importer import ImportErrorMessage, import_file
from services.pdf_export import build_pdf
from services.text_export import to_markdown, to_text
from services.word_export import build_docx

RICH = """# Main

Para with **bold** and \\* a star.

- a
  - nested
    - deep
- b

1. one
  1. sub one
  2. sub two
2. two

| Name | Qty |
|---|---|
| Apple | **3** |
| Pear |

---
End"""


def shape(content):
    return [(b.kind, b.level, b.number) for b in parse_blocks(content)]


def test_nested_lists_numbering_tables_rules():
    assert shape(RICH) == [
        ("heading", 1, 0), ("paragraph", 0, 0), ("bullet", 0, 0), ("bullet", 1, 0), ("bullet", 2, 0),
        ("bullet", 0, 0), ("number", 0, 1), ("number", 1, 1), ("number", 1, 2), ("number", 0, 2),
        ("table", 0, 0), ("rule", 0, 0), ("paragraph", 0, 0)]
    table = next(b for b in parse_blocks(RICH) if b.kind == "table")
    assert len(table.rows) == 3 and all(len(r) == 2 for r in table.rows)      # separator dropped, padded
    assert table.rows[1][1] == [("3", True, False)]


def test_literal_asterisk_and_math_not_italic():
    block = parse_blocks("Para \\* star and 2 * 3 * 4")[0]
    assert "".join(t for t, _, _ in block.lines[0]) == "Para * star and 2 * 3 * 4"
    assert all(not b and not i for _, b, i in block.lines[0])


def test_list_numbering_restarts_after_paragraph():
    nums = [b.number for b in parse_blocks("1. a\n2. b\n\ntext\n\n1. c") if b.kind == "number"]
    assert nums == [1, 2, 1]


def test_stats():
    assert count_stats("") == {"words": 0, "characters": 0, "minutes": 0}
    stats = count_stats("# Title here\n\n**two** words " + "word " * 400)
    assert stats["words"] == 2 + 2 + 400 and stats["minutes"] == 3


def test_preview_markdown_nested_table_rule():
    md = preview_markdown(RICH)
    assert "    - deep" in md and "| Name | Qty |" in md and "| --- | --- |" in md and "\n---\n" in md
    assert "<" not in preview_markdown("<b>x</b>").replace("\\<", "")


def test_docx_contains_tables_nested_lists_rule_and_page_numbers():
    data = build_docx("Rich", RICH, page_size="Letter", page_numbers=True)
    doc = Document(BytesIO(data))
    assert len(doc.tables) == 1
    assert [c.text for c in doc.tables[0].rows[0].cells] == ["Name", "Qty"]
    assert doc.tables[0].rows[0].cells[0].paragraphs[0].runs[0].bold
    styles = [p.style.name for p in doc.paragraphs]
    assert {"List Bullet", "List Bullet 2", "List Bullet 3"} <= set(styles)
    assert abs(doc.sections[0].page_width.inches - 8.5) < 0.01
    with zipfile.ZipFile(BytesIO(data)) as z:
        assert any(b"PAGE" in z.read(n) for n in z.namelist() if "footer" in n)
        assert b"w:pBdr" in z.read("word/document.xml")


def test_docx_a4_and_no_page_numbers():
    data = build_docx("A4", "x", page_size="A4", page_numbers=False)
    doc = Document(BytesIO(data))
    assert abs(doc.sections[0].page_width.inches - 8.27) < 0.01
    with zipfile.ZipFile(BytesIO(data)) as z:
        assert not any(b"PAGE" in z.read(n) for n in z.namelist() if "footer" in n)


def test_pdf_with_tables_and_sizes():
    for size in ("A4", "Letter"):
        assert build_pdf("Rich", RICH, "d", page_size=size).startswith(b"%PDF-")
    wide = "| " + " | ".join(f"H{i}" for i in range(9)) + " |\n" + "\n".join(
        "| " + " | ".join("word " * 8 for _ in range(9)) + " |" for _ in range(40))
    assert build_pdf("Wide table", wide, page_numbers=False).startswith(b"%PDF-")


def test_pdf_uses_bundled_unicode_font_for_non_ascii():
    from services import pdf_export
    pdf_export._fonts.cache_clear()
    _, unicode_ok = pdf_export._fonts()
    assert unicode_ok
    assert (pdf_export._FONT_DIR / "DejaVuSans.ttf").exists()


def test_text_exports():
    md = to_markdown("My Title", "# Heading\n\ntext", "A description")
    assert md.startswith("# My Title\n\n*A description*\n\n# Heading") and md.endswith("\n")
    assert "description" not in to_markdown("T", "x", "description", include_description=False)
    txt = to_text("My Title", RICH, "desc")
    assert txt.startswith("My Title\n========\n\ndesc") and "\u2022 a" in txt and "    \u2022 deep" in txt
    assert "2. two" in txt and "Apple | 3" in txt and "-" * 40 in txt
    assert to_text("", "").startswith("Untitled document")


def test_docx_roundtrip_through_importer():
    result = import_file("rich.docx", build_docx("Rich", RICH))
    assert result["title"] == "Rich"
    assert shape(result["content"]) == shape(RICH)
    assert "**3**" in result["content"]


def test_import_text_and_markdown():
    assert import_file("notes.txt", "Café notes\nline".encode("utf-8")) == {
        "title": "notes", "content": "Café notes\nline"}
    md = import_file("report.md", b"# Real Title\n\nBody")
    assert md["title"] == "Real Title" and md["content"].startswith("# Real Title")
    assert import_file("legacy.txt", "caf\xe9".encode("cp1252"))["content"] == "caf\xe9"
    assert import_file("bom.txt", b"\xef\xbb\xbfhello")["content"] == "hello"
    assert import_file("u16.txt", "hi caf\xe9".encode("utf-16"))["content"] == "hi caf\xe9"


@pytest.mark.parametrize("name,data", [("x.pdf", b"%PDF"), ("empty.txt", b""),
                                       ("bad.docx", b"not a zip"), ("big.txt", b"x" * (5 * 1024 * 1024 + 1))])
def test_import_rejects_bad_files(name, data):
    with pytest.raises(ImportErrorMessage):
        import_file(name, data)


def test_import_docx_native_word_list_numbering():
    doc = Document()
    doc.add_paragraph("First", style="List Number")
    doc.add_paragraph("Bullet", style="List Bullet")
    buf = BytesIO()
    doc.save(buf)
    content = import_file("w.docx", buf.getvalue())["content"]
    assert "1. First" in content and "- Bullet" in content
