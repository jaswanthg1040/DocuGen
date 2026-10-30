"""Generate valid .docx files with python-docx."""
from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from services.formatting import clean_text, parse_blocks

EMPTY_NOTICE = "This document is empty."
PAGE_SIZES = {"A4": (8.27, 11.69), "Letter": (8.5, 11.0)}
_THEME_ATTRS = ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme")
_BULLET_STYLES = ("List Bullet", "List Bullet 2", "List Bullet 3")


def _set_style_font(style, name: str) -> None:
    style.font.name = name
    rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in _THEME_ATTRS:           # theme fonts would override our font
        rfonts.attrib.pop(qn(attr), None)
    rfonts.set(qn("w:eastAsia"), name)


def _add_runs(paragraph, runs, size=None, color=None, force_bold=False) -> None:
    for text, bold, italic in runs:
        run = paragraph.add_run(text)
        run.bold = True if (bold or force_bold) else None
        run.italic = True if italic else None
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor(*color)


def _add_lines(paragraph, lines) -> None:
    for index, runs in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        _add_runs(paragraph, runs)


def _add_field(run, instruction: str) -> None:
    for kind, text in (("begin", None), (None, instruction), ("end", None)):
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        else:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        run._r.append(element)


def _add_page_numbers(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    _add_field(run, " PAGE ")


def _add_rule(doc) -> None:
    paragraph = doc.add_paragraph()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for attr, value in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "A5B4FC")):
        bottom.set(qn(f"w:{attr}"), value)
    border.append(bottom)
    paragraph._p.get_or_add_pPr().append(border)


def _add_table(doc, rows) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, cell_runs in enumerate(row):
            cell = table.cell(r, c)
            _add_runs(cell.paragraphs[0], cell_runs, force_bold=(r == 0))
            if r == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:val"), "clear")
                shade.set(qn("w:color"), "auto")
                shade.set(qn("w:fill"), "E0E7FF")
                cell._tc.get_or_add_tcPr().append(shade)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build_docx(title: str, content: str, description: str = "", font: str = "Calibri",
               include_description: bool = True, page_size: str = "A4",
               page_numbers: bool = True) -> bytes:
    """Return the bytes of a Word document for the given title and content."""
    title = clean_text(title).strip() or "Untitled document"
    description = clean_text(description).strip()
    font = font or "Calibri"
    width, height = PAGE_SIZES.get(page_size, PAGE_SIZES["A4"])

    doc = Document()
    for section in doc.sections:
        section.page_width, section.page_height = Inches(width), Inches(height)
        section.left_margin = section.right_margin = Inches(1)
        section.top_margin = section.bottom_margin = Inches(1)
        if page_numbers:
            _add_page_numbers(section)
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", *_BULLET_STYLES):
        _set_style_font(doc.styles[style_name], font)
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Normal"].paragraph_format.space_after = Pt(8)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.15
    doc.core_properties.title = title[:250]
    doc.core_properties.author = "DocuGen"

    doc.add_paragraph(style="Title").add_run(title)
    if include_description and description:
        _add_runs(doc.add_paragraph(), [(description, False, True)], size=12, color=(0x55, 0x55, 0x66))

    blocks = parse_blocks(content)
    if not blocks:
        _add_runs(doc.add_paragraph(), [(EMPTY_NOTICE, False, True)], color=(0x88, 0x88, 0x88))

    for block in blocks:
        if block.kind == "heading":
            _add_runs(doc.add_paragraph(style=f"Heading {block.level}"), block.lines[0])
        elif block.kind == "bullet":
            _add_runs(doc.add_paragraph(style=_BULLET_STYLES[block.level]), block.lines[0])
        elif block.kind == "number":
            # Typed numbers + hanging indent: every list restarts at 1, unlike Word's
            # shared "List Number" counter.
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.5 + 0.25 * block.level)
            para.paragraph_format.first_line_indent = Inches(-0.25)
            para.paragraph_format.space_after = Pt(4)
            para.add_run(f"{block.number}.\t")
            _add_runs(para, block.lines[0])
        elif block.kind == "rule":
            _add_rule(doc)
        elif block.kind == "table":
            _add_table(doc, block.rows)
        else:
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            _add_lines(para, block.lines)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
